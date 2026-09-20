"""Builds the SAMPLE customer tracking workbook (fictional customer C001).
Run: python3 build_workbook.py  — then recalculate (see ../README.md) so cached values exist.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from datetime import date, datetime
from pathlib import Path
OUT = Path(__file__).resolve().parent.parent / 'onehome_customer_tracker_SAMPLE.xlsx'

F='Arial'
BLUE=Font(name=F,size=10,color='0000FF')
BLK=Font(name=F,size=10,color='000000')
BOLD=Font(name=F,size=10,bold=True)
HDR=Font(name=F,size=10,bold=True,color='FFFFFF')
HFILL=PatternFill('solid',fgColor='0F172A')
GFILL=PatternFill('solid',fgColor='E5E7EB')   # formula cells
SFILL=PatternFill('solid',fgColor='DBEAFE')   # group rows
thin=Side(style='thin',color='CBD5E1'); BRD=Border(left=thin,right=thin,top=thin,bottom=thin)
CEN=Alignment(horizontal='center',vertical='center',wrap_text=True)
RGT=Alignment(horizontal='right',vertical='center',wrap_text=True)
DATE='DD/MM/YYYY'; DT='DD/MM/YYYY HH:MM'; ILS='#,##0 "₪"'

wb=Workbook()
ws_leg=wb.active; ws_leg.title='הסבר'
ws_c=wb.create_sheet('לקוחות'); ws_k=wb.create_sheet('Checklist')
ws_h=wb.create_sheet('שעות'); ws_i=wb.create_sheet('פניות'); ws_p=wb.create_sheet('חבילות')
for ws in wb.worksheets: ws.sheet_view.rightToLeft=True

def header(ws,row,names,widths=None):
    for i,n in enumerate(names,1):
        c=ws.cell(row=row,column=i,value=n); c.font=HDR; c.fill=HFILL; c.alignment=CEN; c.border=BRD
    ws.row_dimensions[row].height=34
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[L(i)].width=w

# ---------------- חבילות ----------------
header(ws_p,1,['חבילה','מחיר התקנה (₪)','ריטיינר חודשי (₪)'],[18,18,18])
for r,(n,a,b) in enumerate([('EcoHome',5700,160),('SafeHome',8700,220),('EliteHome',13150,350)],2):
    ws_p.cell(r,1,n).font=BLK
    for col,v in ((2,a),(3,b)):
        c=ws_p.cell(r,col,v); c.font=BLUE; c.number_format=ILS
    for col in (1,2,3): ws_p.cell(r,col).border=BRD
ws_p['A6']='מקור המחירים: PRD 1 / תמחור מעודכן ב-16/09/2026. אם התמחור משתנה — לעדכן כאן בלבד (תאים כחולים).'
ws_p['A6'].font=Font(name=F,size=9,italic=True,color='475569')
ws_p['A8']='סיכום'; ws_p['A8'].font=BOLD
summ=[('לקוחות פעילים',"=COUNTIF('לקוחות'!$L$2:$L$21,\"פעיל\")",'0'),
      ('הכנסה חודשית מריטיינר (MRR)',"=SUMIFS('לקוחות'!$G$2:$G$21,'לקוחות'!$L$2:$L$21,\"פעיל\")",ILS),
      ('סה"כ פניות פתוחות',"=COUNTIF('פניות'!$F$2:$F$500,\"פתוחה\")",'0'),
      ('סה"כ שעות התקנה שנרשמו',"=SUM('שעות'!$D$2:$D$500)",'0.0')]
for k,(lab,fm,nf) in enumerate(summ,9):
    ws_p.cell(k,1,lab).font=BLK
    c=ws_p.cell(k,2,fm); c.font=BLK; c.fill=GFILL; c.number_format=nf; c.border=BRD
ws_p.column_dimensions['A'].width=34

# ---------------- לקוחות ----------------
cols=['מזהה לקוח','שם לקוח','טלפון','עיר / אזור','חבילה','מחיר התקנה','ריטיינר חודשי',
      'תאריך סקר בית','תאריך התקנה','תאריך אישור קבלה','תחילת ריטיינר (= אישור קבלה)','סטטוס',
      'IP קבוע','MAC של ה-HA','גרסת HA','Nabu Casa','שירותי Nabu Casa שנבחרו',
      'שם bucket ב-R2','environment ב-Sentry','עבודת חשמל','חשמלאי',
      'התקדמות checklist','סה"כ שעות התקנה','הכנסת התקנה לשעה (לפני חומרה)','פניות פתוחות',
      'הפניה לסודות (מנהל סיסמאות)','הערות']
widths=[10,22,14,14,12,12,12,13,13,14,16,14,14,20,11,10,24,24,18,11,14,13,12,16,10,30,30]
header(ws_c,1,cols,widths)
formula_cols={6,7,11,22,23,24,25}
for r in range(2,22):
    ws_c.cell(r,6).value=f'=IF($E{r}="","",IFERROR(INDEX(\'חבילות\'!$B$2:$B$4,MATCH($E{r},\'חבילות\'!$A$2:$A$4,0)),""))'
    ws_c.cell(r,7).value=f'=IF($E{r}="","",IFERROR(INDEX(\'חבילות\'!$C$2:$C$4,MATCH($E{r},\'חבילות\'!$A$2:$A$4,0)),""))'
    ws_c.cell(r,11).value=f'=IF($J{r}="","",$J{r})'
    ws_c.cell(r,22).value=f'=IF($A{r}="","",IFERROR(INDEX(\'Checklist\'!$C$30:$G$30,MATCH($A{r},\'Checklist\'!$C$1:$G$1,0)),""))'
    ws_c.cell(r,23).value=f'=IF($A{r}="","",SUMIFS(\'שעות\'!$D$2:$D$500,\'שעות\'!$A$2:$A$500,$A{r}))'
    ws_c.cell(r,24).value=f'=IF($A{r}="","",IFERROR($F{r}/$W{r},""))'
    ws_c.cell(r,25).value=f'=IF($A{r}="","",COUNTIFS(\'פניות\'!$B$2:$B$500,$A{r},\'פניות\'!$F$2:$F$500,"פתוחה"))'
    for c in range(1,len(cols)+1):
        cell=ws_c.cell(r,c); cell.border=BRD; cell.alignment=RGT
        if c in formula_cols: cell.font=BLK; cell.fill=GFILL
        else: cell.font=BLUE
    for c in (6,7,24): ws_c.cell(r,c).number_format=ILS
    for c in (8,9,10,11): ws_c.cell(r,c).number_format=DATE
    ws_c.cell(r,22).number_format='0%'; ws_c.cell(r,23).number_format='0.0'
sample=['C001','לקוח לדוגמה (לא אמיתי)','050-0000000','תל אביב','SafeHome',None,None,
        date(2026,10,5),date(2026,10,19),date(2026,10,19),None,'פעיל',
        '192.168.1.50','AA:BB:CC:DD:EE:FF','2026.9.3','לא','—',
        'onehome-c001-backups','c001-sample','כן','חשמלאי א׳ (לדוגמה)',
        None,None,None,None,'Bitwarden: OneHome/C001 (לדוגמה)','נתוני דוגמה בלבד — למחוק את השורה לפני שימוש אמיתי']
for c,v in enumerate(sample,1):
    if v is not None: ws_c.cell(2,c,v)
ws_c.freeze_panes='C2'
# validations
def dv(formula,rng):
    d=DataValidation(type='list',formula1=formula,allow_blank=True); ws_c.add_data_validation(d); d.add(rng)
dv("='חבילות'!$A$2:$A$4",'E2:E21')
dv('"ליד,סקר בית,הצעה נשלחה,הזמנת ציוד,התקנה,בדיקות קבלה,פעיל,מושהה,הופסק"','L2:L21')
dv('"כן,לא"','P2:P21'); dv('"כן,לא"','T2:T21')
ws_c['V1'].comment=Comment('מחושב אוטומטית מלשונית Checklist. מופיע רק ללקוחות בשורות 2–6 (5 הלקוחות הראשונים).','Claude')

# ---------------- Checklist ----------------
items=[
('רשת וגישה מקומית',['AP/Client Isolation כבוי בראוטר','רשת שטוחה אחת (או VLAN מאושר ללקוח מיוחד)','DHCP Reservation לפי MAC בראוטר','Internal URL מוגדר לפי ה-IP הקבוע','אפליקציית Companion מזהה את ה-SSID הביתי','גישת אדמין לראוטר אומתה לפני ההתקנה']),
('Nabu Casa (רק אם הלקוח בחר)',['הלקוח פתח חשבון בשמו ובכרטיס שלו','External URL מוגדר','שירותים נבחרו בנפרד (Alexa/Google, גיבוי ענן, Webhooks)','TTS/STT ענני — החלטה מוצהרת (ברירת מחדל: מקומי)','WebRTC למצלמות — החלטה מוצהרת (ברירת מחדל: כבוי)']),
('פרטיות וגיבויים',['מסמך פרטיות נחתם','bucket מבודד ב-R2 נוצר','סיסמת הצפנה ייחודית נשמרה במנהל סיסמאות','גיבוי ראשון הצליח ל-R2','עותק גיבוי מקומי נשמר בבית']),
('בדיקות קבלה — "הבית חי"',['כל מכשירי החבילה מחוברים, בשמות ברורים ופועלים','בדיקת ניתוק אינטרנט: אוטומציה מקומית ממשיכה לרוץ','Push בדיקה התקבל בטלפון הלקוח מחוץ ל-WiFi הביתי','חיישן הצפה נבדק (מים על החיישן מפעילים התראה)','גלאי עשן וסירנה מקומית נבדקו','כניסה מרחוק מדאטה סלולרי (אם נרכשה גישה מרחוק)','הלקוח ביצע בעצמו פעולה אחת (אור / פתיחת קריאה)','ניטור Alert פועל','IP, MAC וגרסת HA נרשמו בלשונית לקוחות']),
('חוזי ומסירה',['אישור כתוב מהחשמלאי (אם הייתה עבודת חשמל)','חבילת מסירה נמסרה ללקוח','אישור קבלה חתום — מתחיל ריטיינר ו-SLA']),
]
header(ws_k,1,['קבוצה','פריט','','','','',''],[26,58,12,12,12,12,12])
for j in range(5):
    c=ws_k.cell(1,3+j,f"=IF('לקוחות'!$A{2+j}=\"\",\"\",'לקוחות'!$A{2+j})")
    c.font=HDR; c.fill=HFILL; c.alignment=CEN; c.border=BRD
r=2; first=2
sample_vals={}
for grp,its in items:
    for it in its:
        ws_k.cell(r,1,grp).font=Font(name=F,size=9,color='64748B'); ws_k.cell(r,2,it).font=BLK
        for c in range(1,8):
            ws_k.cell(r,c).border=BRD; ws_k.cell(r,c).alignment=RGT if c<3 else CEN
        for c in range(3,8): ws_k.cell(r,c).font=BLUE
        # sample customer C001: Nabu Casa group not relevant, remote test not relevant
        v='✓'
        if grp.startswith('Nabu Casa') or 'דאטה סלולרי' in it: v='לא רלוונטי'
        ws_k.cell(r,3,v)
        r+=1
last=r-1
assert last==29, last
ws_k.cell(30,2,'התקדמות (לא כולל "לא רלוונטי")').font=BOLD
for c in range(3,8):
    col=L(c)
    cell=ws_k.cell(30,c,f'=IF({col}$1="","",COUNTIF({col}2:{col}29,"✓")/MAX(1,ROWS({col}2:{col}29)-COUNTIF({col}2:{col}29,"לא רלוונטי")))')
    cell.font=BLK; cell.fill=GFILL; cell.number_format='0%'; cell.alignment=CEN; cell.border=BRD
ws_k.cell(30,1).border=BRD; ws_k.cell(30,2).border=BRD
ws_k.freeze_panes='C2'
d=DataValidation(type='list',formula1='"✓,✗,לא רלוונטי"',allow_blank=True); ws_k.add_data_validation(d); d.add('C2:G29')
ws_k.conditional_formatting.add('C2:G29',CellIsRule(operator='equal',formula=['"✓"'],fill=PatternFill('solid',bgColor='D1FAE5')))
ws_k.conditional_formatting.add('C2:G29',CellIsRule(operator='equal',formula=['"✗"'],fill=PatternFill('solid',bgColor='FEE2E2')))
ws_k.conditional_formatting.add('C2:G29',CellIsRule(operator='equal',formula=['"לא רלוונטי"'],fill=PatternFill('solid',bgColor='E5E7EB')))

# ---------------- שעות ----------------
header(ws_h,1,['מזהה לקוח','שלב','תאריך','שעות בפועל','הערה'],[12,26,13,13,44])
phases=['פגישת מכירה','סקר בית','הצעה וחתימה','הכנה אצל גיא','עבודת חשמל (ליווי)','התקנה ורשת','בדיקות קבלה','מסירה והדרכה','מעקב אחרי מסירה']
rows=[('C001','פגישת מכירה',date(2026,9,28),1.5,'פגישה ראשונה בבית הלקוח'),
      ('C001','סקר בית',date(2026,10,5),2,''),
      ('C001','הצעה וחתימה',date(2026,10,8),1.5,'כולל מסמך פרטיות'),
      ('C001','הכנה אצל גיא',date(2026,10,14),3,'HA, R2, גיבוי ראשון'),
      ('C001','התקנה ורשת',date(2026,10,19),5,''),
      ('C001','בדיקות קבלה',date(2026,10,19),1.5,''),
      ('C001','מסירה והדרכה',date(2026,10,19),1,'')]
for r in range(2,102):
    for c in range(1,6):
        cell=ws_h.cell(r,c); cell.font=BLUE; cell.border=BRD; cell.alignment=RGT
    ws_h.cell(r,3).number_format=DATE; ws_h.cell(r,4).number_format='0.0'
for i,row in enumerate(rows,2):
    for c,v in enumerate(row,1): ws_h.cell(i,c,v)
d=DataValidation(type='list',formula1='"'+','.join(phases)+'"',allow_blank=True); ws_h.add_data_validation(d); d.add('B2:B101')
ws_h.freeze_panes='A2'
ws_h['G1']='שורה לכל שלב עבודה. הסכום לכל לקוח מוצג בלשונית לקוחות.'; ws_h['G1'].font=Font(name=F,size=9,italic=True,color='475569')
ws_h.column_dimensions['G'].width=50

# ---------------- פניות ----------------
header(ws_i,1,['מזהה פנייה','מזהה לקוח','תאריך פתיחה','נושא','חומרה','סטטוס','תאריך פתרון','זמן פתרון (שעות)','הערות'],[12,12,17,44,12,12,17,14,40])
for r in range(2,102):
    for c in range(1,10):
        cell=ws_i.cell(r,c); cell.border=BRD; cell.alignment=RGT
        cell.font=BLK if c==8 else BLUE
        if c==8: cell.fill=GFILL; cell.number_format='0.0'
    ws_i.cell(r,3).number_format=DT; ws_i.cell(r,7).number_format=DT
    ws_i.cell(r,8).value=f'=IF(OR($C{r}="",$G{r}=""),"",ROUND(($G{r}-$C{r})*24,1))'
ws_i.append  # noqa
samp=[('P001','C001',datetime(2026,10,21,9,10),'התראת הצפה לא הגיעה לטלפון של בת הזוג','קריטית','סגורה',datetime(2026,10,21,11,40),'נתוני דוגמה — הוגדרה מחדש הרשאת התראות באפליקציה'),
      ('P002','C001',datetime(2026,11,2,18,0),'שאלה על הוספת מנורה לדשבורד','רגילה','פתוחה',None,'נתוני דוגמה')]
for i,row in enumerate(samp,2):
    ws_i.cell(i,1,row[0]);ws_i.cell(i,2,row[1]);ws_i.cell(i,3,row[2]);ws_i.cell(i,4,row[3]);ws_i.cell(i,5,row[4]);ws_i.cell(i,6,row[5]);ws_i.cell(i,7,row[6]);ws_i.cell(i,9,row[7])
d=DataValidation(type='list',formula1='"קריטית,רגילה"',allow_blank=True); ws_i.add_data_validation(d); d.add('E2:E101')
d=DataValidation(type='list',formula1='"פתוחה,סגורה"',allow_blank=True); ws_i.add_data_validation(d); d.add('F2:F101')
ws_i.freeze_panes='A2'

# ---------------- הסבר ----------------
ws_leg.column_dimensions['A'].width=110
lines=[('OneHome — גיליון מעקב לקוחות (קובץ דוגמה)',True),
('הקובץ הזה מכיל נתוני דוגמה בלבד (לקוח C001, לא אמיתי). כדי להשתמש בו: לשמור עותק במקום פרטי, למחוק את שורת הדוגמה בכל לשונית, ולהתחיל למלא.',False),
('',False),
('איך לקרוא את הצבעים',True),
('טקסט כחול = תא שממלאים ידנית.   טקסט שחור על רקע אפור = נוסחה, לא לערוך.',False),
('',False),
('הלשוניות',True),
('לקוחות — שורה לכל לקוח. נוסחאות מוכנות עד שורה 21. שדות ה-IP, MAC, גרסת HA, R2 ו-Sentry הם אלה ש-PRD 4 (כלי ניהול פנימי) יצטרך, כך שבהמשך אפשר לייבא אותם.',False),
('Checklist — עמודה לכל לקוח, נמשכת אוטומטית מלשונית לקוחות (5 הלקוחות הראשונים). מבוסס על PRD 3, סעיפים B ו-D. "לא רלוונטי" לא נספר באחוז ההתקדמות.',False),
('שעות — שורה לכל שלב עבודה. זה המדד "שעות בפועל לכל שלב" מ-PRD 3, והבסיס לחישוב הכנסה לשעה.',False),
('פניות — יומן פניות אחרי המסירה. רמות החומרה כאן (קריטית/רגילה) הן הצעה — להתאים לרמות ה-SLA המדויקות ב-PRD 1.',False),
('חבילות — מחירי ההתקנה והריטיינר (מקור: PRD 1, 16/09/2026) וסיכום קצר: לקוחות פעילים, הכנסה חודשית וסה"כ פניות פתוחות.',False),
('',False),
('כללים חשובים',True),
('1. אין סודות בקובץ. סיסמאות WiFi, סיסמאות אדמין של HA וסיסמת הצפנת הגיבויים — רק במנהל סיסמאות. בעמודה "הפניה לסודות" כותבים איפה הם שמורים.',False),
('2. לא לשמור את הקובץ האמיתי ב-repo של smart-home או ב-plan-site (האתר ב-GitHub Pages עלול להיות נגיש ברשת) ולא בחשבון OneDrive של המעסיק.',False),
('3. "הכנסת התקנה לשעה" היא מחיר ההתקנה חלקי שעות העבודה, לפני עלות החומרה והחשמלאי — זה לא רווח.',False),
('4. תחילת ריטיינר = תאריך אישור קבלה חתום (החלטה ב-PRD 3, סעיף B).',False)]
for i,(t,b) in enumerate(lines,1):
    c=ws_leg.cell(i,1,t); c.font=Font(name=F,size=12 if i==1 else 10,bold=b); c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='right')
wb.save(OUT)
print('saved',OUT)
