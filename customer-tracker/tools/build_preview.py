"""Builds preview.html (mobile view of the workbook) from the recalculated xlsx.
Run after build_workbook.py + recalculation.
"""
import json
from pathlib import Path
HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
from datetime import datetime, date
from openpyxl import load_workbook
wb=load_workbook(ROOT/'onehome_customer_tracker_SAMPLE.xlsx',data_only=True)
def fmt(v,kind=None):
    if v is None or v=='': return ''
    if isinstance(v,datetime):
        return v.strftime('%d/%m/%Y %H:%M') if (v.hour or v.minute) else v.strftime('%d/%m/%Y')
    if isinstance(v,date): return v.strftime('%d/%m/%Y')
    return v
def money(v): return '' if v in (None,'') else f'{v:,.0f} ₪'
c=wb['לקוחות']; H=[x.value for x in c[1]]
custs=[]
for r in range(2,22):
    row=[c.cell(r,i).value for i in range(1,len(H)+1)]
    if row[0]: custs.append(row)
def cell(row,i):
    v=row[i]
    if i in (5,6,23): return money(v)   # F,G,X (0-based 5,6,23)
    if i==21: return '' if v in (None,'') else f'{v*100:.0f}%'
    if i==22: return '' if v in (None,'') else f'{v:.1f}'
    return fmt(v)
cust_out=[[cell(r,i) for i in range(len(H))] for r in custs]
k=wb['Checklist']; ids=[k.cell(1,i).value for i in range(3,8)]
nc=len([x for x in ids if x])
items=[{'g':k.cell(r,1).value,'t':k.cell(r,2).value,'v':[k.cell(r,3+j).value or '' for j in range(nc)]} for r in range(2,30)]
prog=[k.cell(30,3+j).value for j in range(nc)]
h=wb['שעות']; hours=[[fmt(h.cell(r,i).value) for i in range(1,6)] for r in range(2,102) if h.cell(r,1).value]
i_=wb['פניות']; issues=[[fmt(i_.cell(r,j).value) for j in range(1,10)] for r in range(2,102) if i_.cell(r,1).value]
p=wb['חבילות']; pk=[[p.cell(r,1).value,money(p.cell(r,2).value),money(p.cell(r,3).value)] for r in range(2,5)]
sm=[[p.cell(r,1).value, (money(p.cell(r,2).value) if 'MRR' in p.cell(r,1).value else fmt(p.cell(r,2).value))] for r in range(9,13)]
leg=[wb['הסבר'].cell(r,1).value for r in range(1,20)]
leg=[x for x in leg]
data={'H':H,'cust':cust_out,'ids':ids[:nc],'items':items,'prog':[f'{x*100:.0f}%' for x in prog],'hours':hours,'issues':issues,'pk':pk,'sm':sm,'leg':leg}
tpl=open(HERE/'preview_template.html',encoding='utf-8').read()
open(ROOT/'preview.html','w',encoding='utf-8').write(tpl.replace('/*DATA*/null',json.dumps(data,ensure_ascii=False)))
print(len(cust_out),nc,len(items),len(hours),len(issues))
